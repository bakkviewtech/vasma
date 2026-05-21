from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = "VASMA System Training Manual - Swahili.docx"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            tbl_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "CBD5E1")


def width(cell, value):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(value))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width(cell, widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(rows, widths):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    borders(t)
    set_table_width(t, widths)
    for r_idx, row_data in enumerate(rows):
        for c_idx, value in enumerate(row_data):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(value))
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            if r_idx == 0:
                run.bold = True
                shade(cell, HEADER_FILL)
            p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()


def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = BLUE if level < 3 else DARK_BLUE


def para(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = INK


def bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        r = p.add_run(item)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)


def steps(items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        r = p.add_run(item)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)


def callout(title, body):
    add_table([[title], [body]], [9360])
    t = doc.tables[-1]
    shade(t.rows[0].cells[0], HEADER_FILL)
    shade(t.rows[1].cells[0], LIGHT_FILL)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

for style_name, size, color in (
    ("Normal", 11, INK),
    ("Title", 24, INK),
    ("Heading 1", 16, BLUE),
    ("Heading 2", 13, BLUE),
    ("Heading 3", 12, DARK_BLUE),
):
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = color

header = section.header.paragraphs[0]
header.text = "VASMA System - Mwongozo wa Mafunzo"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(9)
header.runs[0].font.color.rgb = MUTED

title = doc.add_paragraph()
title.style = doc.styles["Title"]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("VASMA System App").bold = True
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Mwongozo wa Mafunzo na Matumizi")
r.font.size = Pt(14)
r.font.color.rgb = DARK_BLUE

callout("Lengo la Mafunzo", "Mtumiaji aweze kusajili wateja, kuunda job card, kuthibitisha kazi, kuchapisha bill/invoice, kurekodi malipo, kuona receipt, kusimamia expenses na commissions, kutoa reports, na kufanya backup.")

heading("1. Utangulizi wa App", 1)
para("VASMA System ni app ya kusimamia huduma za magari. Inafanya kazi kwenye browser, inaweza kusakinishwa kama PWA, na inahifadhi taarifa kwenye browser local storage.")
add_table([
    ["Module", "Kazi yake"],
    ["Dashboard", "Kuonyesha muhtasari wa biashara, job cards, billed, paid, outstanding, expenses, na commissions."],
    ["Customers", "Kusajili mteja, namba ya gari, vehicle model, contact person, na mobile number."],
    ["Job Cards", "Kuunda rekodi ya kazi, huduma, attendant, action, purchased items, notes, na reference."],
    ["Work Confirmation", "Kuchagua job card, kuthibitisha kazi, ku-preview bill/invoice, na ku-export PDF."],
    ["Payments", "Kurekodi malipo, kuchuja receipts, kuona receipt, na ku-export receipt."],
    ["Admin Edit", "Sehemu moja ya ku-edit au delete transactional records."],
    ["Settings", "Business profile, language, theme, page setup, install, backup, na restore."],
], [1900, 7460])

heading("2. Kuanzisha App", 1)
steps([
    "Fungua folder la VASMA System.",
    "Run Start-VASMA.bat.",
    "Fungua http://127.0.0.1:4173/ kwenye browser.",
    "Kama unaona version ya zamani, refresh au clear cache ya site.",
])

heading("3. Mpangilio wa Lugha", 1)
steps([
    "Fungua Settings.",
    "Nenda kwenye Language.",
    "Chagua English au Swahili.",
    "Bonyeza Save Settings.",
])
para("Ukiweka Swahili, sehemu kuu za interface zitaonekana kwa Kiswahili. Data ulizoingiza hazitabadilishwa.")

heading("4. Usajili wa Wateja", 1)
steps([
    "Fungua Customers.",
    "Bonyeza Add Customer.",
    "Jaza Customer Name, Truck/Vehicle Number, Vehicle Model, Mobile Number, na Contact Person.",
    "Bonyeza Save Customer.",
])

heading("5. Job Cards", 1)
para("Kila job card inapata reference ya kipekee kama JC-000001.")
steps([
    "Fungua Job Cards.",
    "Bonyeza New Job Card.",
    "Chagua tarehe na customer.",
    "Ongeza service item: category, service item, attendant, action, na amount.",
    "Ongeza purchased items kama kuna vifaa vilivyonunuliwa.",
    "Tumia Add Note kama kuna maelekezo ya ziada.",
    "Bonyeza Save Job Card.",
])
bullets([
    "Actions za kawaida ni Checked, Changed, Replaced, na Repaired.",
    "Actions mpya zinaongezwa kwenye Master Data.",
    "Job Card inaweza ku-exportiwa PDF kupitia Export Job Card.",
])

heading("6. Work Confirmation na Bill/Invoice", 1)
steps([
    "Fungua Work Confirmation.",
    "Chagua Job Card Number kwenye dropdown.",
    "Hakiki huduma na vifaa vilivyonunuliwa.",
    "Chagua action na tick Done kwa kazi iliyothibitishwa.",
    "Bonyeza Save Confirmation.",
    "Bonyeza Preview Bill au Export PDF.",
])
callout("QR Code", "Bill/Invoice ina QR code chini ya document. QR ina taarifa muhimu kama invoice number, job card, tarehe, customer, gari, amount, na payment status.")

heading("7. Payments na Receipts", 1)
steps([
    "Fungua Payments.",
    "Bonyeza Record Payment.",
    "Chagua job card yenye outstanding balance.",
    "Chagua Cash, Mobile Money, au Bank.",
    "Weka comment/reference ya SMS au bank.",
    "Weka amount kwa outstanding bill lines tu.",
    "Bonyeza Save Payment.",
])
bullets([
    "Bill iliyolipwa yote haiwezi kupokea payment nyingine.",
    "Bill iliyolipwa sehemu itaonyesha outstanding lines tu.",
    "Receipt reference ni RCPT-JC000001.",
    "Tumia filter kutafuta receipt kwa customer name au vehicle number.",
    "Bonyeza View Receipt ku-preview na Export Receipt ku-save PDF.",
])

heading("8. Expenses", 1)
steps([
    "Fungua Expenses.",
    "Bonyeza Add Expense.",
    "Chagua expense item au ongeza mpya.",
    "Weka amount na comment.",
    "Bonyeza Save Expense.",
])

heading("9. Commissions", 1)
steps([
    "Fungua Commissions.",
    "Bonyeza Allocate Commission.",
    "Chagua job card na attendant.",
    "Chagua Amount au Percent.",
    "Weka rate/amount na uhifadhi.",
])
para("Analysis ya Attendant Income hutumia commission amount ya attendant.")

heading("10. Reports na Analysis", 1)
add_table([
    ["Report", "Maelezo"],
    ["Outstanding Report", "Inaonyesha bill amount, paid amount, na outstanding."],
    ["Income Statement per Job Card", "Inaonyesha paid amount, commission, expenses, na balance."],
    ["Commission Report", "Inaonyesha commission kwa attendant na job card."],
], [2600, 6760])
steps([
    "Fungua Reports.",
    "Chagua From na To date.",
    "Chagua report type.",
    "Export PDF au Export Excel.",
])
bullets([
    "Service Category Income hupanga category kwa income.",
    "Service Item Income hupanga huduma kwa income.",
    "Attendant Income hutumia commissions.",
    "Date Income hupanga tarehe kwa income.",
])

heading("11. Master Data", 1)
para("Master Data hutunza dropdown values zinazotumika kwenye app.")
add_table([
    ["Master Data", "Matumizi"],
    ["Employees", "Attendants na supervisors."],
    ["Packages", "Car wash packages."],
    ["Categories", "Makundi ya huduma."],
    ["Service Items", "Huduma ndani ya category."],
    ["Expense Items", "Aina za expenses."],
    ["Actions", "Checked, Changed, Replaced, Repaired, na nyingine."],
], [2300, 7060])

heading("12. Admin Edit", 1)
steps([
    "Fungua Admin Edit.",
    "Chagua record type.",
    "Chagua record husika.",
    "Bonyeza Edit au Delete.",
])
callout("Tahadhari", "Ukifuta job card, malipo na commissions zilizounganishwa nayo zinaweza kuondolewa. Tumia Admin Edit kwa umakini.")

heading("13. Settings", 1)
bullets([
    "Business Setup: business name, location, owner name, mobile.",
    "User Profile: jina la user hutumika kwenye signature ya reports, bills, job cards, na receipts.",
    "Language: English au Swahili.",
    "Theme: light au dark mode.",
    "Page Setup: Job Card, Bill/Invoice, na Receipt layout.",
    "App Tools: Install, Export Backup, Import Backup.",
])
heading("13.1 Page Setup", 2)
bullets([
    "Fit one page - Compact: chaguo linalopendekezwa.",
    "Normal: nafasi kubwa zaidi kwa documents fupi.",
    "Maximum fit - Dense: inasaidia document ndefu ku-fit ukurasa mmoja.",
])
heading("13.2 Backup", 2)
steps([
    "Fungua Settings.",
    "Bonyeza Export Backup ku-save JSON file.",
    "Hifadhi file mahali salama.",
    "Kurudisha data, bonyeza Import Backup na chagua JSON file.",
])

heading("14. Printing na PDF", 1)
bullets([
    "Export PDF hutumia browser print dialog.",
    "Chagua Save as PDF kuunda PDF.",
    "QR code iko chini ya invoice/bill na receipt.",
    "Chagua Compact au Dense kwenye Settings kama document haitoshi ukurasa mmoja.",
])

heading("15. Reference Numbers", 1)
add_table([
    ["Document", "Format", "Mfano"],
    ["Job Card", "JC-000001", "JC-000001"],
    ["Invoice", "INV-JC000001", "INV-JC000001"],
    ["Receipt", "RCPT-JC000001", "RCPT-JC000001"],
], [2200, 3000, 4160])

heading("16. Troubleshooting", 1)
add_table([
    ["Tatizo", "Sababu", "Suluhisho"],
    ["App inaonyesha folder listing", "Server imeanzishwa kwenye folder tofauti.", "Run Start-VASMA.bat ndani ya VASMA System folder."],
    ["Version ya zamani inaonekana", "Cache ya browser/PWA.", "Refresh au clear site data."],
    ["Payment haikubaliwi", "Bill imelipwa yote.", "Angalia outstanding balance."],
    ["QR haionekani", "Internet haipo kwa QR image service.", "Unganisha internet na reload preview."],
    ["PDF haifit page moja", "Layout ni normal au lines ni nyingi.", "Chagua Maximum fit - Dense kwenye Settings."],
], [2100, 3300, 3960])

heading("17. Checklist ya Mkufunzi", 1)
bullets([
    "Onyesha kusajili customer na vehicle model.",
    "Tengeneza job card yenye service items na purchased items.",
    "Fanya work confirmation na export bill/invoice.",
    "Rekodi partial payment na onyesha outstanding lines.",
    "View na export receipt.",
    "Ongeza expense na commission.",
    "Generate reports na analysis.",
    "Tumia Admin Edit kurekebisha record.",
    "Badilisha language English/Swahili.",
    "Export na import backup.",
])

doc.save(OUT)
print(OUT)
