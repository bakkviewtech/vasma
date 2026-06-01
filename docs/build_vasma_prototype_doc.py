from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE


OUT = "docs/VASMA_System_Comprehensive_Prototype_Specification.docx"


APP_NAME = "VASMA System - Vehicle Auto Service Management System"
PREPARED_FOR = "Viewtech / Bakari Kamanga"
VERSION = "Prototype Specification - Last Update Baseline"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = width
                set_cell_margins(row.cells[idx])
                row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, bold=False, italic=False, size=None, color=None):
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):
            lead, rest = item
            r = p.add_run(lead)
            r.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "VASMA Table"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "E8EEF5")
        cell.text = header
        for p in cell.paragraphs:
            p.runs[0].bold = True
            p.runs[0].font.color.rgb = RGBColor(11, 37, 69)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    if widths:
        set_table_width(table, widths)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()
    return table


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "VASMA Callout"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    style_run(r, bold=True, color="0B2545")
    p.add_run(" " + body)
    doc.add_paragraph()


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)

    table_style = styles.add_style("VASMA Table", WD_STYLE_TYPE.TABLE)
    table_style.font.name = "Calibri"
    table_style.font.size = Pt(9)

    callout_style = styles.add_style("VASMA Callout", WD_STYLE_TYPE.TABLE)
    callout_style.font.name = "Calibri"
    callout_style.font.size = Pt(10)


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = "VASMA System Prototype Specification"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if p.runs:
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(91, 113, 139)
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "Prepared for app builder use - confidential working specification"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if p.runs:
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(91, 113, 139)


def section_title(doc, text):
    doc.add_heading(text, level=1)


def subsection(doc, text):
    doc.add_heading(text, level=2)


def subsub(doc, text):
    doc.add_heading(text, level=3)


def build_doc():
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(APP_NAME)
    style_run(r, bold=True, size=22, color="0B2545")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Comprehensive Prototype and Builder Specification")
    style_run(r, size=14, color="2E74B5")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"{VERSION}\nPrepared for: {PREPARED_FOR}\nDeveloper company: Viewtech\nDeveloper profile: Bakari Kamanga, +255 767 528 039")
    add_callout(
        doc,
        "Purpose:",
        "This document is the handover blueprint for rebuilding VASMA System exactly as agreed in the latest update. It describes the product behavior, modules, data model, security, subscription flow, reports, print/PDF output, mobile behavior, and acceptance criteria.",
    )

    section_title(doc, "1. Product Vision and Scope")
    add_paragraph(
        doc,
        "VASMA System is a production-ready vehicle auto service management system for businesses that provide Tyre Service, Car Wash, General Service, or any approved combination of those service categories. The app must run as a web/PWA application on Android, iOS, Windows, Mac, and desktop browsers while using MySQL as the main production database.",
    )
    add_bullets(
        doc,
        [
            ("Primary users: ", "vehicle service businesses, supervisors, attendants, business owners, and the system owner."),
            ("Core outcome: ", "create job cards, confirm work, bill customers, receive payments, allocate commissions, manage expenses, generate reports, and control subscription access."),
            ("Commercial model: ", "multi-business, multi-branch, subscription-based access with package-specific service categories."),
            ("Data principle: ", "MySQL is the main database. Browser storage is only offline cache or temporary pending sync."),
            ("Security principle: ", "access must be enforced by UI, route, API, and database filters, not by menu hiding only."),
        ],
    )

    subsection(doc, "Supported platforms")
    add_table(
        doc,
        ["Platform", "Requirement"],
        [
            ["Android", "Installable PWA; mobile browser layout with hamburger sidebar and touch-friendly controls."],
            ["iOS", "Installable web app behavior through Safari; responsive views and print/PDF-ready documents."],
            ["Windows", "Full desktop browser experience; local or hosted backend; printable documents."],
            ["Mac", "Full desktop browser experience; PWA-capable where browser supports it."],
            ["Web", "HTTPS deployment, backend API, MySQL, JWT, role-based access, subscription rules."],
        ],
        [Inches(1.4), Inches(5.1)],
    )

    section_title(doc, "2. Recommended Tech Stack")
    add_table(
        doc,
        ["Layer", "Specification"],
        [
            ["Frontend", "HTML5, CSS3, JavaScript, PWA manifest, service worker. Current app may remain framework-free unless builder chooses React/Vue while preserving behavior."],
            ["Backend", "Node.js with Express, REST APIs, JWT authentication, bcrypt password hashing, CORS restricted to frontend origin."],
            ["Database", "MySQL as main database with normalized multi-tenant tables and JSON mirror only where migration compatibility is needed."],
            ["Offline cache", "IndexedDB preferred; localStorage acceptable only for small metadata. All offline changes require sync_status = pending/synced/failed."],
            ["Reports", "Excel export with numeric values clean and currency shown in headers, PDF/print preview, QR code on documents."],
            ["Deployment", "Backend on Render/Railway/VPS, MySQL managed or VPS, frontend via same backend or Netlify/Vercel with HTTPS."],
        ],
        [Inches(1.35), Inches(5.15)],
    )

    section_title(doc, "3. Roles and Access")
    add_table(
        doc,
        ["Role", "Purpose", "Key access"],
        [
            ["Super Admin / System Owner", "Viewtech owner panel in a separate folder/link.", "Packages, subscriptions, landing pages, payment methods, business requests, activation, suspension, extension approval."],
            ["Business Owner / Business Admin", "Owns one business account and its branches.", "Business setup, branches, users, permissions, data view, service package use, reports, operational oversight."],
            ["Supervisor", "Operational manager inside a business/branch.", "Job cards, confirmations, payments, reports, commission follow-up according to permissions."],
            ["User", "Normal business operator.", "Transactions allowed by role and subscription package."],
            ["Developer", "Special developer mode for controlled UI/layout adjustments.", "Login as Developer with approved credentials; edit headings/layout sizing and save configuration without corrupting business data."],
        ],
        [Inches(1.4), Inches(2.15), Inches(2.95)],
    )

    section_title(doc, "4. Login, Trial, Subscription, and Activation")
    subsection(doc, "Login screens")
    add_bullets(
        doc,
        [
            "Business app login must accept demo credentials before activation: username demo and password demo123.",
            "Admin/business owner credentials continue to work according to their assigned roles.",
            "Super Admin must use a separate admin link and folder, not the normal business app login.",
            "Logout icon must be a small red power symbol: ⏻.",
        ],
    )
    subsection(doc, "Business code format")
    add_table(
        doc,
        ["Business Code", "Service Code", "Business No.", "Services Enabled", "Initial PIN Example"],
        [
            ["VASMA-1-01", "1", "01", "Tyre Service only", "1010"],
            ["VASMA-2-02", "2", "02", "Car Wash only", "2020"],
            ["VASMA-3-03", "3", "03", "General Service only", "3030"],
            ["VASMA-12-04", "12", "04", "Tyre Service + Car Wash", "1204"],
            ["VASMA-13-05", "13", "05", "Tyre Service + General Service", "1305"],
            ["VASMA-23-06", "23", "06", "Car Wash + General Service", "2306"],
            ["VASMA-123-07", "123", "07", "Full Package", "1230"],
        ],
        [Inches(1.2), Inches(0.9), Inches(0.85), Inches(2.65), Inches(0.9)],
    )
    add_bullets(
        doc,
        [
            "If service code contains 1, enable Tyre Service.",
            "If service code contains 2, enable Car Wash.",
            "If service code contains 3, enable General Service.",
            "Services not included in the code must be hidden and blocked at API/database level.",
            "The Business Number identifies the tenant/business account.",
        ],
    )

    subsection(doc, "Demo trial flow")
    add_numbers(
        doc,
        [
            "User logs in with demo / demo123.",
            "App displays service category multiple-choice options: Car Wash, Tyre Service, General Service, any combination of two, or all service categories.",
            "User saves selected services.",
            "App opens Business Setup / Business Profile. User enters business name, owner name, location, mobile number, and new password.",
            "Trial runs for 14 days and only selected service modules are visible and usable.",
            "During trial, app saves all transactions with record_mode = TRIAL and archive flag = false.",
            "When trial expires, new transactions are restricted. Reports remain visible. User can request extension or subscription.",
        ],
    )

    subsection(doc, "Subscription request flow")
    add_numbers(
        doc,
        [
            "Business owner opens Subscription page.",
            "Business owner chooses package/services and submits subscription request.",
            "App displays available package options and prices.",
            "Request goes to Super Admin for review.",
            "Super Admin approves, edits, upgrades, downgrades, or rejects the request.",
            "If approved, status becomes APPROVED_AWAITING_PAYMENT.",
            "Business owner submits payment details in the subscription payment page.",
            "Super Admin confirms payment and activates business with service package, branch limit, period, start date, and expiry date.",
            "Business receives business code and OTP/initial PIN.",
            "On first official login, user must change password before continuing.",
        ],
    )

    subsection(doc, "Trial data handling after official activation")
    add_bullets(
        doc,
        [
            ("Continue trial records: ", "trial transactions remain active and become part of official business records."),
            ("Archive trial records: ", "trial records are hidden from official operational views but preserved in archive. This is not a delete/reset."),
            ("Data View setting: ", "Business Setup must provide Official Records, Trial / Archived Records, and All Records views."),
            ("Deletion rule: ", "trial or archived data cannot be deleted/reset unless system owner approves a formal system order requested by business owner."),
        ],
    )

    subsection(doc, "Offline license and extension behavior")
    add_bullets(
        doc,
        [
            "The app must generate encrypted request tokens offline when subscription extension or activation is needed.",
            "Request token must include device ID, business code, business number, service package, requested period, app version, installation date, and current license status.",
            "User sends the token to Super Admin.",
            "Super Admin enters the token on another device and generates an activation/extension code.",
            "Activation code must be tied to device ID and cannot be reused on another device.",
            "Expired account is report-only: no new job cards, payments, expenses, confirmations, or commission allocations.",
        ],
    )

    section_title(doc, "5. Super Admin Panel")
    add_paragraph(
        doc,
        "The Super Admin panel must live in a separate frontend folder and separate URL, for example /superadmin/. It is for the system owner, not for business owners. It manages the commercial platform and subscription lifecycle.",
    )
    add_table(
        doc,
        ["Super Admin module", "Required capability"],
        [
            ["Dashboard", "Overview of businesses, active subscriptions, pending subscription requests, pending payments, trial extensions, expired/suspended accounts."],
            ["Manage Packages", "Create, edit, deactivate, and price packages by service combination, branch limit, user limit, duration, and plan name."],
            ["Manage Subscriptions", "View requests submitted by business owners; approve, edit, reject, suspend, cancel, expire, extend, upgrade, or downgrade."],
            ["Registered Businesses", "Review businesses that registered from landing page; inspect owner, contacts, package request, status, branches, and history."],
            ["Subscription Payments", "Review payment details submitted by businesses; confirm or reject; activation starts only after payment confirmation."],
            ["Payment Methods", "Manage payment instructions such as bank, mobile money, reference rules, and admin notes."],
            ["Landing Pages", "Create/update marketing/registration pages and package presentation shown to potential business owners."],
            ["Activation Tools", "Decode request/extension/password-reset tokens and generate offline activation codes."],
            ["Audit Logs", "Record all subscription, business status, package, payment, and activation actions."],
        ],
        [Inches(1.65), Inches(4.85)],
    )

    section_title(doc, "6. Business App Navigation and UI")
    subsection(doc, "Global layout")
    add_bullets(
        doc,
        [
            "Sidebar deep navy #022b4f and active menu bright blue #1e9af0.",
            "Sidebar must be scrollable and long enough to align visually with dashboard graph area.",
            "Mobile must use hamburger sidebar overlay; desktop must keep full sidebar.",
            "Business name must appear at the top across the entire app, centered below the VASMA brand line and above each page heading.",
            "Top-right shows logged-in user and small red logout power icon.",
            "Floating bottom-right Dashboard button remains for quick return unless page-specific action conflicts.",
        ],
    )
    subsection(doc, "Main menu")
    add_table(
        doc,
        ["Menu item", "Submenus / behavior"],
        [
            ["Dashboard", "Default landing page after login. Shows selected business and allowed services."],
            ["Customers", "Customer and vehicle registration with active/inactive status."],
            ["Job Cards", "Create, view, preview, export job cards."],
            ["Work Confirmation", "Confirm done items/subitems and generate bill/invoice."],
            ["Payments", "Receive and allocate payments, view/export receipts."],
            ["Service Cards", "Create General Service card based on job card items."],
            ["Expenses", "Record expense category, free-entry expense item, amount, and comment."],
            ["Commissions", "Allocate commissions only for unprocessed items and alert pending commission."],
            ["Transaction Edit", "Central edit/delete area for transactional records."],
            ["Business Admin", "Business profile, branches, users, permissions, audit logs."],
            ["Master Data", "Employees/attendants, packages, categories, service items, service subitems, expense categories, actions."],
            ["Reports", "Date-filtered reports, Excel/PDF exports."],
            ["Analysis", "Dashboard charts and visual analysis."],
            ["License", "License status, request code, activation code, extension token, records."],
            ["Security", "Security settings, sessions, permissions, audit."],
            ["Settings", "Business setup, theme, language, page setup, app tools, about app."],
        ],
        [Inches(1.55), Inches(4.95)],
    )

    section_title(doc, "7. Dashboard Prototype")
    add_paragraph(
        doc,
        "The dashboard should follow a clean management dashboard style: compact summary cards at the top, one main chart at a time, tabbed content, and a right-side insight/ranking panel. It must not look like a crowded table/report page.",
    )
    add_table(
        doc,
        ["Dashboard area", "Behavior"],
        [
            ["Header", "Brand line centered: VASMA SYSTEM - VEHICLE AUTO SERVICE MANAGEMENT SYSTEM. Business name centered below. Page title Dashboard below."],
            ["Period tabs", "Horizontal tabs: Today, MTD, YTD, Custom. Content updates according to selected period."],
            ["Summary cards", "Job Cards, Billed, Paid, Outstanding. Cards are compact with one key metric and a small indicator/chip."],
            ["Second-row tabs", "Overview, Job Status, Service Performance, Finance, Commission. Only selected tab content is visible."],
            ["Main chart card", "One large chart only. Chart changes by selected second-row tab and period."],
            ["Right insight card", "Shows top service items, pending confirmations, payment status, or commission pending depending on selected tab."],
            ["Mobile", "Cards stack, chart appears before insight card, tables scroll horizontally, touch controls remain large."],
        ],
        [Inches(1.55), Inches(4.95)],
    )
    subsection(doc, "Dashboard charts")
    add_bullets(
        doc,
        [
            "Job Card Status Summary vertical bar chart with labels on top of bars.",
            "Service Category performance bar chart using ratio-to-highest scaling so bars are not too tall.",
            "Service Item performance bar chart using ratio-to-highest scaling.",
            "Attendant/Commission performance bar chart using ratio-to-highest scaling.",
            "Amount labels use compact K formatting where appropriate: 10,000 = 10K, 15,500 = 15.5K.",
            "No floating graph behavior. Graph size can be configured through developer edit mode only.",
        ],
    )

    section_title(doc, "8. Operational Modules")
    subsection(doc, "Customers")
    add_bullets(
        doc,
        [
            "Register customer name, contact person, mobile number, vehicle number, vehicle model, and optional notes.",
            "Customer status Active/Inactive. Inactive customers must not appear in new job cards but must remain on old records.",
            "Vehicle model examples: Mitsubishi Pickup, Toyota Corolla, Toyota Land Cruiser V8.",
        ],
    )
    subsection(doc, "Job Cards")
    add_bullets(
        doc,
        [
            "Job Card reference must be sequential and unique in format JC-000001.",
            "Invoice reference generated from job card: INV-JC000001.",
            "Receipt reference generated from job card: RCPT-JC000001.",
            "Top line fields: Date, Customer, Vehicle Model, Mileage.",
            "Additional service lines keep heading only at the top; extra rows show fields only.",
            "Action column is removed from Job Card entry page.",
            "Service line fields: Category, Service Item, Attendant, Amount.",
            "Each service item may have associated subitems shown as multiple-choice checkboxes inside the item.",
            "Purchased items section links job card number, category, and service item. Table headings: Item name, unit, Qty, Unit cost, Amount.",
            "Note field is hidden by default and appears when user clicks Add note.",
            "Paid job cards must be locked from normal editing and only editable through Transaction Edit with proper role.",
        ],
    )
    subsection(doc, "Work Confirmation and Bill/Invoice")
    add_bullets(
        doc,
        [
            "Filter by Job Card Number dropdown, not free recall input.",
            "After Attendant column, show Job Status with item/subitem checkboxes from job card.",
            "User ticks done service subitems and can add free remarks/comment at the end.",
            "Confirmed work status must appear on preview, invoice/bill, PDF, and print.",
            "If work is confirmed and commission is not allocated, show popup alert with job card number.",
            "Even if payment is complete, if commission is not allocated the alert still appears.",
            "Once job card is paid and commission allocated, it should not appear as pending in confirmation/commission alerts.",
        ],
    )
    subsection(doc, "Payments")
    add_bullets(
        doc,
        [
            "Payment cannot be posted to a bill that is already fully paid.",
            "For partially paid bills, system shows only outstanding amount/items.",
            "Filter should be Excel-like header filters with search/dropdown, not a large separate filter input.",
            "View the receipt option displays payment receipt.",
            "Receipt can be exported/printed/shared.",
            "Paid amount cannot exceed job card amount unless admin confirms overpayment.",
        ],
    )
    subsection(doc, "Expenses")
    add_bullets(
        doc,
        [
            "Expense form uses Expense Category dropdown and Expense Item free-entry field.",
            "Fields: Date, Expense Category, Expense Item, Amount, Comment.",
            "Expense categories are maintained in Master Data and can be active/inactive.",
        ],
    )
    subsection(doc, "Commissions")
    add_bullets(
        doc,
        [
            "Commission allocation filters only records that have not been processed.",
            "Commission cannot exceed paid amount.",
            "Analysis wording: Attendant Income uses Attendant Commission.",
            "Dashboard/alerts show commission pending until allocated, even when bill is paid.",
        ],
    )
    subsection(doc, "Service Cards")
    add_bullets(
        doc,
        [
            "Create General Service card from General Service job card items.",
            "Business name centered at top; customer details placed to the right of 'General Service Card' to save space.",
            "Columns: Service Item, small checkbox, Remarks.",
            "Remove Done, Action, and Attendant columns.",
            "Service Mileage, Next Service Mileage, and Next Service Date appear horizontally.",
            "Maintain Service Interval comment.",
            "Bottom signatures: Attended by on left, Reviewed by on right, name and signature.",
            "Card must fit appropriate standard page size for print/PDF.",
        ],
    )

    section_title(doc, "9. Master Data and Settings")
    subsection(doc, "Master Data")
    add_bullets(
        doc,
        [
            "Master records display as selectable dropdown lists with Add, Edit, Delete, and Status where applicable.",
            "Employees/attendants have Active/Inactive status. Inactive attendants do not appear in new job cards.",
            "Categories have Active/Inactive status. Inactive categories do not appear in new job cards.",
            "Service item belongs to a category.",
            "Service subitem belongs to a service item and appears as multiple-choice checkboxes during transactions.",
            "Deleted master data must not affect old transaction records. Store display names/snapshots on transactions.",
        ],
    )
    subsection(doc, "Business Admin")
    add_bullets(
        doc,
        [
            "Business Profile merges old User Profile and Business Setup details.",
            "Fields: Business Name, owner/user name, location, mobile number, optional email, branch defaults, password change.",
            "Business name from profile appears on dashboard, reports, job cards, invoices, receipts, exported PDFs, and about page.",
            "Username/profile name appears on report and bill signatures.",
            "Branches page creates/edits branches under the business within subscription branch limit.",
            "Users and Roles page creates staff accounts and assigns roles/branch access.",
            "Permissions page controls module/action permissions.",
            "Audit Logs page shows important actions and security events.",
        ],
    )
    subsection(doc, "Settings")
    add_bullets(
        doc,
        [
            "Settings sections use compact droplist/button style: Business Profile, Theme, Language, Page Setup, App Tools, About the App.",
            "Language options: English and Swahili. App should support multi-language labels where possible.",
            "Theme options: Light and Dark.",
            "Page Setup: Job Card, Bill/Invoice, Receipt sizing to fit one page before PDF export.",
            "App Tools: Install, Export Backup, Import Backup, database/sync tools.",
            "About the App must show App Name, Business Name, Platform, App Developer, Developer Company Name, Developer Mobile, Data Storage.",
            "Developer Company Name: Viewtech.",
        ],
    )

    section_title(doc, "10. Reports, Exports, Printing, and QR Codes")
    add_bullets(
        doc,
        [
            "Date format across app: dd-mmm-yyyy, for example 19-May-2026.",
            "Report date filter format: From: dd-mmm-yyyy to: dd-mmm-yyyy horizontally.",
            "Report creator/signature must come from logged-in user profile.",
            "Excel exports must put currency in amount column headers, for example Amount (Tsh), Paid Amount (Tsh), Commission (Tsh), Expenses (Tsh), Balance (Tsh).",
            "Excel export values must be clean numbers without Tsh prefix so Excel analysis works.",
            "All summations must be correct for Paid Amount, Commission, Expenses, and Balance.",
            "Payment report headers should match payment page structure.",
            "Print preview must open before print/save as PDF where supported by browser.",
        ],
    )
    subsection(doc, "QR code rule")
    add_bullets(
        doc,
        [
            "Barcodes are removed everywhere. Use QR code only.",
            "QR code appears professionally at the bottom of invoice, bill, receipt, print preview, downloaded PDF, shared document, and printed hard copy.",
            "QR code should be small enough to keep document on one page but clear and high contrast.",
            "QR content should contain a secure verification link or key document details: number, date, customer, vehicle, total amount, payment status.",
            "Label examples: Scan to verify invoice or Scan to view bill details.",
        ],
    )

    section_title(doc, "11. Database Specification")
    add_paragraph(
        doc,
        "The database must be multi-business, multi-branch, subscription-aware, and service-package-aware. Every business-owned transactional record must carry business_id. Branch-specific records must also carry branch_id. Offline-originated records must carry sync_status.",
    )
    add_table(
        doc,
        ["Table", "Purpose / important fields"],
        [
            ["businesses", "Business account: business_code, business_name, owner_name, mobile, location, status, created_at."],
            ["branches", "Branch records: business_id, branch_name, location, mobile, status."],
            ["users", "Global/user login records: username, password_hash, role, business_id, branch_id, must_change_password, status."],
            ["roles", "Role definitions such as super_admin, business_admin, supervisor, user, developer."],
            ["service_categories", "Tyre Service, Car Wash, General Service and business/custom categories."],
            ["service_items", "Items linked to category: Basic Wash, Full Wash, Puncture Repair, General Service item, etc."],
            ["service_sub_items", "Multiple-choice subitems linked to service_item_id."],
            ["subscription_plans", "Package name, service_code, price, period, branch limit, user limit, status."],
            ["business_subscriptions", "Business subscription lifecycle: status, start_date, expiry_date, service_code, branch_limit."],
            ["subscription_services", "Allowed service categories per subscription."],
            ["subscription_payments", "Business submitted subscription payment details and Super Admin confirmation."],
            ["customers", "Customer and vehicle details, status, business_id."],
            ["job_cards", "JC ref, customer, vehicle, mileage, totals, status, record_mode, is_archived, sync_status."],
            ["job_card_items", "Job card service lines with category/item/attendant/amount and display snapshots."],
            ["job_card_item_sub_items", "Selected service subitems for each job card item."],
            ["purchased_items", "Purchased material lines linked to job card/category/service item."],
            ["invoices", "Invoice/bill generated from job card with INV-JC000001 ref, totals, QR data."],
            ["payments", "Payment allocations linked to invoice/job card, outstanding calculations, sync status."],
            ["receipts", "Receipt refs RCPT-JC000001, payment details, QR data."],
            ["expenses", "Expense category, free-entry item, amount, comments, business/branch."],
            ["commissions", "Attendant commission allocations and processed status."],
            ["service_cards", "General Service card header, mileage, next service details, signatures."],
            ["service_card_items", "Service card lines with checkbox and remarks."],
            ["activation_requests", "Offline request/extension/activation tokens and status."],
            ["password_reset_requests", "Offline reset token and admin reset-code workflow."],
            ["audit_logs", "User, module, action, before/after summary, IP/device, timestamp."],
            ["payment_methods", "Super Admin managed payment method instructions."],
            ["landing_pages", "Super Admin managed landing page content/package display."],
            ["vasma_data", "Compatibility JSON snapshots only; not the primary long-term normalized store."],
            ["sync_audit", "Offline sync attempts, success/failure, error message."],
            ["app_settings", "Per-business settings: theme, language, page sizes, data view, developer layout config."],
        ],
        [Inches(1.85), Inches(4.65)],
    )

    subsection(doc, "Tenant and subscription rules")
    add_bullets(
        doc,
        [
            "Business cannot see data from another business.",
            "Branch users cannot see another branch unless role grants multi-branch access.",
            "Service category cannot be used if subscription does not allow it.",
            "Expired/suspended account is report-only.",
            "All transactions must save to MySQL and optionally mirror to offline cache.",
            "All writes from offline cache must sync with sync_status pending, synced, or failed.",
        ],
    )

    section_title(doc, "12. API Specification")
    add_table(
        doc,
        ["Endpoint group", "Required endpoints / behavior"],
        [
            ["Auth", "/api/auth/login, /api/auth/business-login, /api/auth/change-password, /api/auth/logout."],
            ["Public/Landing", "/api/public/packages, /api/public/register-business, landing content endpoints."],
            ["Business context", "/api/data/context loads business, branch, subscription, services, settings, user permissions."],
            ["Data sync", "/api/data/save and /api/data/load for local-to-MySQL compatibility; normalized endpoints preferred for production."],
            ["Subscriptions", "/api/data/subscription-options, /api/data/subscription-request, /api/data/subscription-payment, /api/data/trial-extension-request."],
            ["Super Admin", "/api/admin/packages, /api/admin/businesses, /api/admin/subscriptions, /api/admin/subscription-payments, /api/admin/payment-methods, /api/admin/landing-pages."],
            ["Activation", "/api/license/request, /api/license/activate, offline token decode/generate endpoints in admin context."],
            ["Reports", "Date-filtered reports, Excel/PDF metadata, business/branch/service filters."],
        ],
        [Inches(1.55), Inches(4.95)],
    )

    section_title(doc, "13. Validation Rules")
    add_table(
        doc,
        ["Rule", "Expected behavior"],
        [
            ["Job Card Ref unique", "System must generate sequential unique JC-000001 references and reject duplicates."],
            ["Vehicle/Truck number required", "Customer/job card cannot be saved without vehicle registration."],
            ["Contact person required", "Customer/job card requires contact person."],
            ["Mobile number required", "Customer/business setup requires mobile number."],
            ["Service category required", "Each job card line requires allowed active category."],
            ["Service item required", "Each job card line requires active item under selected category."],
            ["Amount numeric", "Allow comma display formatting but store numeric clean value."],
            ["Paid amount limit", "Paid amount cannot exceed job card/invoice amount unless admin confirms overpayment."],
            ["Commission limit", "Commission cannot exceed paid amount."],
            ["Outstanding", "Outstanding = Job Card Amount - Paid Amount."],
            ["Balance", "Balance = Paid Amount - Commission - Expenses."],
            ["Old records safe", "Deleted/inactive master data must not break old records."],
            ["Inactive attendants", "Do not show inactive attendants in new job cards."],
            ["Paid lock", "Paid job cards are locked from normal edit; use Transaction Edit with permission."],
        ],
        [Inches(1.75), Inches(4.75)],
    )

    section_title(doc, "14. Security Requirements")
    add_bullets(
        doc,
        [
            "Passwords stored using bcrypt hashes only.",
            "JWT tokens expire; recommended default 8 hours or less.",
            "CORS restricted to production frontend URL.",
            "Never commit .env or secrets to GitHub.",
            "Super Admin, Business Admin, and Developer access must be separate and auditable.",
            "API must enforce business_id, branch_id, role, subscription status, and service code checks on every protected request.",
            "Audit log must record login, failed login, password changes, transaction edits/deletes, subscription changes, activation, extension, and payment confirmation.",
            "Password reset uses offline reset token: business owner generates reset token, sends to Super Admin, Super Admin generates reset code, user enters reset code and creates new password.",
        ],
    )

    section_title(doc, "15. Mobile and Responsive Requirements")
    add_bullets(
        doc,
        [
            "Test widths: 360px, 390px, 412px, 768px, and desktop.",
            "Sidebar becomes collapsible hamburger on mobile.",
            "Dashboard cards stack vertically on mobile.",
            "Tables must not break screen; use horizontal scrolling wrappers.",
            "Buttons must be large enough for touch screens.",
            "Forms use full-width inputs on mobile.",
            "Invoice, bill, receipt, and report pages must display properly on mobile and print/PDF.",
            "QR codes must remain clear and scannable on mobile and printed PDF.",
            "Settings/About tables must fit mobile using wrapping or horizontal scroll without clipping.",
        ],
    )

    section_title(doc, "16. Document References and Number Formats")
    add_table(
        doc,
        ["Document", "Reference format", "Notes"],
        [
            ["Job Card", "JC-000001", "Sequential per business. Must never repeat."],
            ["Invoice/Bill", "INV-JC000001", "Generated from related job card number."],
            ["Receipt", "RCPT-JC000001", "Generated from related payment/job card."],
            ["Payment Reference", "RCPT-JC000001 or payment-specific ref", "Must be searchable in payments and reports."],
            ["Date", "dd-mmm-yyyy", "Example: 19-May-2026 across app and reports."],
            ["Currency display", "TSh 188,000.00", "UI display; Excel exports use numeric values with Tsh in header only."],
        ],
        [Inches(1.45), Inches(1.7), Inches(3.35)],
    )

    section_title(doc, "17. Acceptance Checklist for App Builder")
    add_bullets(
        doc,
        [
            "Login works for demo trial, business users, business admin, developer mode, and separate Super Admin panel.",
            "Demo trial lets user select services and restricts modules accordingly for 14 days.",
            "Business setup/profile feeds business name and user signature across dashboard, reports, forms, PDFs, receipts, and invoices.",
            "All current transactions save to MySQL with business_id and branch_id.",
            "Offline cache can queue pending records and sync to MySQL when online.",
            "Subscription request, payment submission, Super Admin confirmation, activation, expiry, and report-only expired behavior work end to end.",
            "Dashboard follows the agreed template style with period tabs, summary cards, one chart at a time, and right insight panel.",
            "Job cards support categories, service items, associated subitems, attendants, mileage, notes, purchased items, and correct reference generation.",
            "Work confirmation displays job card item/subitem checkboxes and remarks; confirmed status appears in previews and exports.",
            "Payment module blocks fully paid bills, supports partial outstanding, shows receipt, and exports receipt.",
            "Commission module filters unprocessed records and alerts pending commission even after payment until allocated.",
            "Reports calculate totals correctly and export Excel with numeric values and Tsh only in headers.",
            "QR code appears on document outputs; barcode is removed everywhere.",
            "All modules are responsive on Android/iOS/desktop and print cleanly.",
            "Inactive/deleted master records do not break old records and do not appear in new transactions.",
            "Security checks exist on frontend, API, route, and database filtering.",
        ],
    )

    section_title(doc, "18. Implementation Notes")
    add_callout(
        doc,
        "Builder note:",
        "The app should be implemented module by module with database migrations and tests after each stage. Do not convert all localStorage logic at once without a sync layer; use a compatibility migration so existing demo/trial records can be imported safely.",
    )
    add_numbers(
        doc,
        [
            "Build and verify database schema first.",
            "Implement auth, business context, subscription checks, and branch/service access filters.",
            "Migrate operational modules to normalized MySQL endpoints.",
            "Add offline cache/pending sync.",
            "Polish dashboard and responsive UI.",
            "Complete reports, exports, printing, and QR verification.",
            "Run full acceptance checklist before handover.",
        ],
    )

    doc.add_page_break()
    section_title(doc, "Appendix A. Example Service Subitems")
    add_table(
        doc,
        ["Category", "Service Item", "Example Subitems"],
        [
            ["Car Wash", "Full Wash", "Exterior wash; interior vacuum; dashboard cleaning; tyre shine; floor mat cleaning."],
            ["Tyre Service", "Puncture Repair", "Remove wheel; inspect puncture; patch/plug; inflate; leak test; refit wheel."],
            ["General Service", "Engine Service", "Engine oil; oil filter; air cleaner; coolant check; brake fluid check; road test."],
        ],
        [Inches(1.25), Inches(1.6), Inches(3.65)],
    )

    section_title(doc, "Appendix B. Subscription Status Lifecycle")
    add_table(
        doc,
        ["Status", "Meaning", "Allowed user behavior"],
        [
            ["PENDING", "Business registration or subscription request submitted.", "Await review; no official activation yet."],
            ["TRIAL", "Demo trial active.", "Use selected service modules for 14 days."],
            ["APPROVED_AWAITING_PAYMENT", "Super Admin approved package but payment not confirmed.", "Can view instructions and submit payment details."],
            ["PAYMENT_SUBMITTED", "Business submitted payment information.", "Await Super Admin confirmation."],
            ["ACTIVE", "Subscription paid/confirmed and active.", "Full access according to package, role, branch, and service permissions."],
            ["EXPIRED", "Subscription ended.", "Reports only; block new transactions."],
            ["SUSPENDED", "Super Admin paused account.", "Reports only or blocked according to policy."],
            ["CANCELLED", "Subscription cancelled.", "Access blocked except admin/export if allowed."],
        ],
        [Inches(1.65), Inches(2.5), Inches(2.35)],
    )

    section_title(doc, "Appendix C. Deployment Checklist")
    add_bullets(
        doc,
        [
            "Create private GitHub repository.",
            "Commit frontend, backend, superadmin, schema, README, and deployment instructions.",
            "Keep node_modules, .env, backend/.env, logs, and database backups out of GitHub.",
            "Install MySQL and create vasma_db using backend/schema.sql.",
            "Configure backend .env: DB_HOST, DB_USER, DB_PASS, DB_NAME, JWT_SECRET, PORT, FRONTEND_URL.",
            "Deploy backend to Render/Railway/VPS with HTTPS.",
            "Deploy frontend with correct API URL and service worker cache version.",
            "Change all default passwords before production.",
            "Create backup/export/import routines for MySQL.",
            "Test PWA install on Android and desktop.",
        ],
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_doc())
