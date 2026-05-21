from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "VASMA System Training Manual.docx"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "CBD5E1")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(rows, widths, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    set_table_borders(t)
    set_table_width(t, widths)
    for r_idx, row_data in enumerate(rows):
        row = t.rows[r_idx]
        for c_idx, value in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(value))
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            if header and r_idx == 0:
                run.bold = True
                run.font.color.rgb = INK
                set_cell_shading(cell, HEADER_FILL)
            p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    return t


def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = BLUE if level < 3 else DARK_BLUE
    return p


def para(text="", bold_start=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_start and text.startswith(bold_start):
        r1 = p.add_run(bold_start)
        r1.bold = True
        r1.font.color.rgb = INK
        r2 = p.add_run(text[len(bold_start):])
        r2.font.color.rgb = INK
    else:
        r = p.add_run(text)
        r.font.color.rgb = INK
    return p


def bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        r = p.add_run(item)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)


def steps(items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        r = p.add_run(item)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)


def callout(title, body):
    t = table([[title], [body]], [9360], header=False)
    set_cell_shading(t.rows[0].cells[0], HEADER_FILL)
    set_cell_shading(t.rows[1].cells[0], LIGHT_FILL)
    for run in t.rows[0].cells[0].paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = DARK_BLUE


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(11)
styles["Normal"].font.color.rgb = INK
for style_name, size, color in (
    ("Title", 24, INK),
    ("Heading 1", 16, BLUE),
    ("Heading 2", 13, BLUE),
    ("Heading 3", 12, DARK_BLUE),
):
    style = styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = True

header = section.header.paragraphs[0]
header.text = "VASMA System - Training Manual"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(9)
header.runs[0].font.color.rgb = MUTED

footer = section.footer.paragraphs[0]
footer.text = "Prepared for VASMA users"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.runs[0].font.size = Pt(9)
footer.runs[0].font.color.rgb = MUTED

title = doc.add_paragraph()
title.style = styles["Title"]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("VASMA System App")
run.bold = True
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("User Training Manual and Operations Guide")
r.font.size = Pt(14)
r.font.color.rgb = DARK_BLUE
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("Covers: setup, job cards, work confirmation, billing, payments, receipts, reports, analysis, admin edit, backup, and page setup")
r.font.size = Pt(10)
r.font.color.rgb = MUTED

callout("Training Goal", "By the end of this guide, a user should be able to register customers, create job cards, confirm work, print bills/invoices, record payments, view receipts, manage expenses and commissions, generate reports, and maintain the application safely.")

heading("1. App Overview", 1)
para("VASMA System is an offline-first web application for vehicle auto service operations. It can run in a browser and can be installed as a Progressive Web App on Android, iOS, Windows, macOS, and web-supported devices.")
table([
    ["Module", "Main purpose"],
    ["Dashboard", "Shows business snapshot, recent job cards, billing, paid amount, outstanding balance, expenses, and commissions."],
    ["Customers", "Registers customer, vehicle number, vehicle model, contact person, and mobile number."],
    ["Job Cards", "Creates service work records with service items, attendants, actions, purchased items, notes, and unique job card references."],
    ["Work Confirmation", "Allows the user to select a job card, confirm work done, preview bill/invoice, and export PDF."],
    ["Payments", "Records payments, filters receipts, views receipts, and exports receipts."],
    ["Expenses", "Records operating expenses."],
    ["Commissions", "Allocates attendant commission by amount or percentage."],
    ["Admin Edit", "Central place for editing and deleting transactional records."],
    ["Master Data", "Maintains employees, packages, categories, service items, expense items, and actions."],
    ["Reports and Analysis", "Generates reports and rankings for business review."],
    ["Settings", "Controls business profile, user profile, theme, page setup, installation, backup, and restore."],
], [1800, 7560])

heading("2. Getting Started", 1)
heading("2.1 Start the app", 2)
steps([
    "Open the VASMA System folder.",
    "Run Start-VASMA.bat, or run python -m http.server 4173 --bind 127.0.0.1 from the app folder.",
    "Open http://127.0.0.1:4173/ in a browser.",
    "If the browser shows an old version, refresh the page or clear site cache because the app is cached for offline use.",
])
heading("2.2 Install on device", 2)
bullets([
    "Android: open the app in Chrome and choose Install app or Add to Home screen.",
    "iOS: open in Safari, tap Share, then Add to Home Screen.",
    "Windows/macOS: open in Edge or Chrome and choose Install app where available.",
    "Installation tools are located in Settings, together with backup and restore.",
])

heading("3. Recommended Daily Workflow", 1)
table([
    ["Step", "User action", "Module"],
    ["1", "Register customer and vehicle model if not already registered.", "Customers"],
    ["2", "Create a job card with service items, actions, attendants, notes, and purchased items.", "Job Cards"],
    ["3", "Confirm work done and print/export the bill or invoice.", "Work Confirmation"],
    ["4", "Record payment only for outstanding bill lines.", "Payments"],
    ["5", "View and export receipt for the payment.", "Payments"],
    ["6", "Allocate attendant commission where applicable.", "Commissions"],
    ["7", "Record expenses and review reports/analysis.", "Expenses, Reports, Analysis"],
], [900, 5580, 2880])

heading("4. Customers Module", 1)
para("Use this module to create customer and vehicle records before preparing job cards.")
steps([
    "Open Customers.",
    "Click Add Customer.",
    "Enter customer name, truck/vehicle number, vehicle model, mobile number, and contact person.",
    "Save the customer.",
])
callout("Important", "Vehicle Model is used on job card and invoice screens. Example models include Mitsubishi Pickup, Toyota Corolla, Toyota Land Cruiser V8, and similar vehicle types.")

heading("5. Job Cards Module", 1)
para("A job card is the main work record. Every job card receives a unique sequential reference in the format JC-000001.")
heading("5.1 Create a job card", 2)
steps([
    "Open Job Cards.",
    "Click New Job Card.",
    "Select date and customer. Vehicle model is displayed from customer registration.",
    "Add service items by selecting category, service item, attendant, action, and amount.",
    "Add purchased items where materials were bought for the job. Enter item name, unit, quantity, unit cost, and amount.",
    "Click Add Note only when additional instructions are needed.",
    "Save Job Card.",
])
heading("5.2 Actions and status", 2)
bullets([
    "Common actions include Checked, Changed, Replaced, and Repaired.",
    "More actions can be added in Master Data.",
    "The action is later confirmed in Work Confirmation.",
])
heading("5.3 Export Job Card", 2)
steps([
    "Open or view the job card.",
    "Click Export Job Card.",
    "Use the browser print dialog to save as PDF or print hard copy.",
    "Page layout is controlled from Settings > Job Card Page Setup.",
])

heading("6. Work Confirmation and Bill/Invoice", 1)
para("Work Confirmation is used before issuing the customer bill/invoice. It uses a dropdown filter, so details appear only after selecting a job card.")
steps([
    "Open Work Confirmation.",
    "Select a Job Card Number from the dropdown.",
    "Review customer, vehicle, services, purchased items, and bill summary.",
    "Confirm service action and tick Done where the work is verified.",
    "Update purchased material status where needed.",
    "Click Save Confirmation.",
    "Click Preview Bill to inspect the bill/invoice.",
    "Click Export PDF to print or save the bill/invoice.",
])
callout("QR Code", "Bills and invoices include a QR code at the bottom of the document. The QR contains invoice details such as invoice number, bill number, date, customer, vehicle, amount, and payment status.")

heading("7. Payments and Receipts", 1)
heading("7.1 Record payment", 2)
steps([
    "Open Payments.",
    "Click Record Payment.",
    "Select a job card with outstanding balance.",
    "Choose payment method: Cash, Mobile Money, or Bank.",
    "Enter SMS or bank reference in the comment field, and attach a screenshot filename if needed.",
    "Pay only the outstanding bill lines shown by the system.",
    "Save payment.",
])
heading("7.2 Payment rules", 2)
bullets([
    "Fully paid bills cannot receive another payment.",
    "Partially paid bills show only outstanding bill lines.",
    "The system prevents payment above the outstanding line amount.",
    "Receipts receive a unique reference in the format RCPT-JC000001.",
])
heading("7.3 View and export receipt", 2)
steps([
    "Use the filter box to search receipts by customer name, vehicle number, vehicle model, job card, method, or comment.",
    "Click View Receipt beside the payment.",
    "Review receipt details and QR code.",
    "Click Export Receipt to print or save PDF.",
])

heading("8. Expenses Module", 1)
steps([
    "Open Expenses.",
    "Click Add Expense.",
    "Select expense item. If the item is not available, choose Add new.",
    "Enter amount and comment.",
    "Save expense.",
])
para("Expense records feed income statement reporting and business balance calculations.")

heading("9. Commissions Module", 1)
steps([
    "Open Commissions.",
    "Click Allocate Commission.",
    "Select date, job card, and attendant.",
    "Choose Amount or Percent.",
    "Enter rate or amount. The system calculates commission amount.",
    "Save commission.",
])
callout("Analysis note", "Attendant Income analysis uses attendant commission totals, not service billing totals.")

heading("10. Reports Module", 1)
para("Reports can be generated by date range. Dates display in dd-Mmm-yyyy format and the selected range appears horizontally as From: dd-Mmm-yyyy to: dd-Mmm-yyyy.")
table([
    ["Report", "Purpose"],
    ["Outstanding Report", "Shows bill amount, paid amount, and outstanding balance by job card."],
    ["Income Statement per Job Card", "Shows paid amount, commission, expenses, and balance by job card."],
    ["Commission Report", "Shows commission paid by attendant and job card."],
], [2500, 6860])
steps([
    "Open Reports.",
    "Choose From and To dates.",
    "Select report type.",
    "Click Export PDF for print/PDF output.",
    "Click Export Excel to download CSV compatible with Excel.",
])

heading("11. Analysis Module", 1)
para("Analysis ranks business results from high to low for the selected period.")
bullets([
    "Service Category Income ranks categories by billed service income.",
    "Service Item Income ranks individual service items by billed income.",
    "Attendant Income ranks attendants by commission amount.",
    "Date Income ranks dates by job card income.",
])

heading("12. Master Data", 1)
para("Master Data controls dropdown values used across the app. Editing master data should be done carefully because it affects future job cards and workflows.")
table([
    ["Master data", "Use"],
    ["Employees", "Attendants and supervisors selected in job cards and commissions."],
    ["Car Wash Packages", "Package names and included services."],
    ["Service Categories", "Top-level service groups such as Tire Service, General Service, and Car Wash."],
    ["Service Items", "Service lines under each category."],
    ["Expense Items", "Expense types used in expense entry."],
    ["Actions", "Work actions such as Checked, Changed, Replaced, Repaired."],
], [2300, 7060])
para("Use the dropdown to select a record, then Edit or Delete. Add buttons remain available for new setup values.")

heading("13. Admin Edit", 1)
para("Admin Edit centralizes edit and delete controls for transactional records. This keeps daily operation screens cleaner and reduces accidental changes.")
steps([
    "Open Admin Edit.",
    "Select record type: Customers, Job Cards, Payments, Expenses, or Commissions.",
    "Select the specific record.",
    "Click Edit to update the record or Delete to remove it.",
])
callout("Caution", "Deleting a job card also removes linked payments and commissions. Use Admin Edit only when correction is necessary.")

heading("14. Settings", 1)
table([
    ["Setting area", "Description"],
    ["Business Setup", "Business name, location, owner name, and mobile number used in documents."],
    ["User Profile", "User name is used as the prepared-by signature on reports, bills, job cards, and receipts."],
    ["Login Setup", "Stores login configuration fields for profile control."],
    ["Theme", "Switches between light and dark mode."],
    ["Page Setup", "Controls Job Card, Bill/Invoice, and Receipt print layout."],
    ["App Tools", "Install app, export backup, and import backup."],
], [2300, 7060])
heading("14.1 Page setup options", 2)
bullets([
    "Fit one page - Compact: recommended default for normal documents.",
    "Normal: wider spacing for fewer line items.",
    "Maximum fit - Dense: smallest practical layout to help long job cards or invoices fit one page.",
])
heading("14.2 Backup and restore", 2)
steps([
    "Open Settings.",
    "Click Export Backup to download JSON backup.",
    "Store backup safely outside the browser.",
    "To restore, click Import Backup and select the JSON file.",
])
callout("Data storage", "The app stores data in browser local storage. Always export backup before clearing browser data, changing computers, or reinstalling the app.")

heading("15. Printing, PDF, and Sharing", 1)
bullets([
    "Export PDF buttons use the browser print dialog. Choose Save as PDF to create a file.",
    "The printed document includes the QR code at the bottom.",
    "For a one-page result, choose Compact or Dense in Settings before exporting.",
    "If the page is still too long, reduce service line count by printing supporting detail separately or use Dense layout.",
])

heading("16. Reference Numbers", 1)
table([
    ["Document", "Format", "Example"],
    ["Job Card", "JC-000001", "JC-000001"],
    ["Invoice", "INV-JC000001", "INV-JC000001"],
    ["Receipt", "RCPT-JC000001", "RCPT-JC000001"],
], [2200, 3000, 4160])

heading("17. Troubleshooting", 1)
table([
    ["Problem", "Likely cause", "Fix"],
    ["App opens folder listing", "Server started in wrong folder.", "Stop server, open VASMA System folder, run Start-VASMA.bat."],
    ["Old screen still appears", "Browser/PWA cache still has old files.", "Refresh page or clear site data; service worker cache updates after reload."],
    ["Cannot record payment", "Bill is fully paid or no outstanding line exists.", "Check Payments and bill outstanding status."],
    ["QR code not visible", "No internet for QR image service.", "Reconnect internet and reload document preview before printing."],
    ["Document does not fit one page", "Too many lines or normal layout selected.", "Use Settings > Page Setup > Maximum fit - Dense."],
], [2100, 3300, 3960])

heading("18. Trainer Checklist", 1)
bullets([
    "Demonstrate customer registration with vehicle model.",
    "Create one job card with service items and purchased items.",
    "Confirm work and export bill/invoice.",
    "Record partial payment and show only outstanding lines on second payment.",
    "View and export receipt.",
    "Add expense and commission.",
    "Generate reports and export PDF/Excel.",
    "Use Admin Edit to correct one sample record.",
    "Set page setup to Dense and re-export a document.",
    "Export and import backup.",
])

doc.save(OUT)
print(OUT)
